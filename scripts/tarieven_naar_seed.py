"""Zet een tarievenlijst (Excel) om naar een dbt-seed in seeds/tarieven/.

Voorbeeld (Wlz 2026):

    poetry run python scripts/tarieven_naar_seed.py "<pad>/tarievenlijst2026.xlsx" \
        --factuurtype "AW319 Wet Langdurige Zorg" \
        --geldig-vanaf 2026-01-01 --geldig-tot 2026-12-31 \
        --uit seeds/tarieven/tarieven_wlz_2026.csv

Het script zoekt de kopregel (de rij met de prestatiecode-kolom), negeert lege
rijen en kolommen, ontdubbelt op prestatiecode en schrijft een standaard-CSV
(komma-gescheiden, punt als decimaalteken, ISO-datums, UTF-8).
Codes zonder tarief (bv. meerzorg, individueel tarief) krijgen een leeg tarief.
"""

import argparse
import csv
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

# Kolomnamen in het bronbestand → kolomnamen in de seed. Wijkt de layout van een
# Wmo- of Jeugdwet-lijst af, voeg dan een eigen mapping toe en kies die via --layout.
LAYOUTS = {
    "wlz": {
        "prestatiecode": "Prestatiecode",
        "prestatienaam": "Prestatienaam",
        "tijdseenheid": "Tijdseenheid",
        "tarief": "Te declareren tarief (per tijdseenheid)",
    },
}

SEED_KOLOMMEN = [
    "factuur_type",
    "debiteurnummer",
    "geldig_vanaf",
    "geldig_tot",
    "prestatiecode",
    "prestatienaam",
    "tijdseenheid",
    "tarief",
]


def lees_tarieven(excel: Path, layout: dict[str, str], tabblad: str | None) -> list[dict]:
    wb = load_workbook(excel, read_only=True, data_only=True)
    tabbladen = [tabblad] if tabblad else [n for n in wb.sheetnames if n.startswith("Tarieven")]
    if not tabbladen:
        sys.exit(f"Geen tabblad 'Tarieven*' gevonden in {excel.name}; geef --tabblad op.")

    rijen = list(wb[tabbladen[0]].iter_rows(values_only=True))
    codekolom = layout["prestatiecode"]
    kop_index = next(
        (i for i, rij in enumerate(rijen) if codekolom in [str(v).strip() if v else v for v in rij]),
        None,
    )
    if kop_index is None:
        sys.exit(f"Kopregel met kolom '{codekolom}' niet gevonden in tabblad '{tabbladen[0]}'.")

    kop = [str(v).strip() if v is not None else None for v in rijen[kop_index]]
    ontbrekend = [bron for bron in layout.values() if bron not in kop]
    if ontbrekend:
        sys.exit(f"Kolommen ontbreken in het bronbestand: {ontbrekend}")
    positie = {doel: kop.index(bron) for doel, bron in layout.items()}

    tarieven = []
    for rij in rijen[kop_index + 1:]:
        if rij[positie["prestatiecode"]] in (None, ""):
            continue
        tarieven.append({doel: rij[i] for doel, i in positie.items()})
    return tarieven


def als_decimaal(waarde) -> Decimal | None:
    return None if waarde in (None, "") else Decimal(str(waarde))


def ontdubbel(tarieven: list[dict]) -> list[dict]:
    """Eén rij per prestatiecode; dubbele codes moeten hetzelfde tarief en dezelfde eenheid hebben."""
    per_code: dict[str, dict] = {}
    conflicten = []
    for t in tarieven:
        code = str(t["prestatiecode"]).strip()
        eerder = per_code.get(code)
        if eerder is None:
            per_code[code] = t
        elif (als_decimaal(eerder["tarief"]), eerder["tijdseenheid"]) != (
            als_decimaal(t["tarief"]),
            t["tijdseenheid"],
        ):
            conflicten.append(code)
    if conflicten:
        sys.exit(f"Dubbele prestatiecodes met verschillend tarief of eenheid: {sorted(set(conflicten))}")
    return list(per_code.values())


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("excel", type=Path)
    parser.add_argument("--factuurtype", required=True, help="Exact zoals factuur_type in mart_facturen")
    parser.add_argument("--geldig-vanaf", required=True, type=date.fromisoformat)
    parser.add_argument("--geldig-tot", required=True, type=date.fromisoformat)
    parser.add_argument("--debiteurnummer", default="", help="Alleen bij tarieven per gemeente/debiteur (Wmo)")
    parser.add_argument("--layout", default="wlz", choices=LAYOUTS)
    parser.add_argument("--tabblad", help="Standaard het eerste tabblad dat met 'Tarieven' begint")
    parser.add_argument("--uit", required=True, type=Path)
    args = parser.parse_args()

    tarieven = lees_tarieven(args.excel, LAYOUTS[args.layout], args.tabblad)
    uniek = ontdubbel(tarieven)

    args.uit.parent.mkdir(parents=True, exist_ok=True)
    with args.uit.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=SEED_KOLOMMEN)
        writer.writeheader()
        for t in uniek:
            writer.writerow({
                "factuur_type": args.factuurtype,
                "debiteurnummer": args.debiteurnummer,
                "geldig_vanaf": args.geldig_vanaf.isoformat(),
                "geldig_tot": args.geldig_tot.isoformat(),
                "prestatiecode": str(t["prestatiecode"]).strip(),
                "prestatienaam": str(t["prestatienaam"] or "").strip(),
                "tijdseenheid": str(t["tijdseenheid"] or "").strip(),
                "tarief": "" if als_decimaal(t["tarief"]) is None else f"{als_decimaal(t['tarief']):f}",
            })

    print(f"{len(tarieven)} rijen gelezen, {len(tarieven) - len(uniek)} dubbel, {len(uniek)} geschreven naar {args.uit}")
    zonder_tarief = [str(t["prestatiecode"]).strip() for t in uniek if als_decimaal(t["tarief"]) is None]
    if zonder_tarief:
        print(f"Zonder tarief (leeg gelaten): {zonder_tarief}")


if __name__ == "__main__":
    main()
